"""
Source adapters. Every fetcher returns a list of rows sorted by date:
  value series : [(date_iso, value), ...]
  ohlc series  : [(date_iso, open, high, low, close, volume), ...]
Raise SourceError on any failure so the caller can try the next candidate.
"""
import csv
import io
import json
import os
import re
import time
from datetime import date, datetime, timezone

import requests

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
TIMEOUT = 25


class SourceError(Exception):
    pass


def _get(url, headers=None, params=None, retries=1, sleep=2.0, timeout=None):
    h = {"User-Agent": UA, "Accept": "*/*"}
    if headers:
        h.update(headers)
    last = None
    for i in range(retries + 1):
        try:
            r = requests.get(url, headers=h, params=params, timeout=timeout or TIMEOUT)
            if r.status_code == 200 and r.content:
                return r
            last = f"HTTP {r.status_code}: {r.text[:120]!r}"
        except requests.RequestException as e:  # noqa: PERF203
            last = repr(e)
        time.sleep(sleep * (i + 1))
    raise SourceError(f"{url} -> {last}")


DATE_RE = re.compile(r"^(\d{4})(?:-(\d{2}))?(?:-(\d{2}))?$")
QUARTER_RE = re.compile(r"^(\d{4})-?Q([1-4])$")
WEEK_RE = re.compile(r"^(\d{4})-W(\d{2})$")


def norm_period(p):
    """Normalise a period string (YYYY, YYYY-MM, YYYY-MM-DD, YYYY-Qn, YYYY-Www) to an ISO date
    at the START of the period. Returns None if unparseable."""
    p = p.strip()
    m = DATE_RE.match(p)
    if m:
        y, mo, d = m.group(1), m.group(2) or "01", m.group(3) or "01"
        return f"{y}-{mo}-{d}"
    m = QUARTER_RE.match(p)
    if m:
        return f"{m.group(1)}-{int(m.group(2)) * 3 - 2:02d}-01"
    m = WEEK_RE.match(p)
    if m:
        try:
            return date.fromisocalendar(int(m.group(1)), int(m.group(2)), 1).isoformat()
        except ValueError:
            return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})T", p)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.match(r"^(\d{2})/(\d{2})/(\d{4})$", p)
    if m:  # dd/mm/yyyy (Bundesbank downloads)
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    return None


def _num(s):
    if s is None:
        return None
    s = str(s).strip().replace(" ", "")
    if s in ("", ".", "NA", "NaN", "null", "-", "n/a", "..."):
        return None
    # Bundesbank/ECB csv can use comma decimals when lang=de; handle both.
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _clean_value_rows(rows):
    out = {}
    for d, v in rows:
        if d and v is not None and v == v:
            out[d] = v
    if not out:
        raise SourceError("no usable observations")
    return sorted(out.items())


# ---------------------------------------------------------------- FRED
FRED_API_KEY = os.environ.get("FRED_API_KEY", "").strip()


def fred(key):
    """FRED series. Uses the official API when FRED_API_KEY is set (reliable from cloud runners),
    otherwise the public fredgraph.csv endpoint, which throttles datacenter IPs."""
    rows = []
    if FRED_API_KEY:
        r = _get("https://api.stlouisfed.org/fred/series/observations",
                 params={"series_id": key, "api_key": FRED_API_KEY, "file_type": "json", "observation_start": "1900-01-01"},
                 headers={"Accept": "application/json"}, retries=1)
        try:
            for o in r.json().get("observations", []):
                d = norm_period(o.get("date", ""))
                if d:
                    rows.append((d, _num(o.get("value"))))
        except ValueError as e:
            raise SourceError(f"fred api bad json for {key}: {e}")
        return _clean_value_rows(rows)
    r = _get("https://fred.stlouisfed.org/graph/fredgraph.csv", params={"id": key}, retries=0, timeout=60)
    for rec in csv.reader(io.StringIO(r.text)):
        if len(rec) < 2:
            continue
        d = norm_period(rec[0])
        if d is None:
            continue
        rows.append((d, _num(rec[1])))
    return _clean_value_rows(rows)


def fred_inv(key):
    return [(d, 1.0 / v) for d, v in fred(key) if v]


def fred_pct(key):
    """Month-over-month % change of a FRED level series."""
    rows = fred(key)
    out = []
    for (d0, v0), (d1, v1) in zip(rows, rows[1:]):
        if v0:
            out.append((d1, round((v1 / v0 - 1.0) * 100.0, 4)))
    return out


def fred_yoy(key):
    rows = fred(key)
    by = dict(rows)
    out = []
    for d, v in rows:
        y, m, dd = d.split("-")
        prev = f"{int(y) - 1}-{m}-{dd}"
        if prev in by and by[prev]:
            out.append((d, round((v / by[prev] - 1.0) * 100.0, 4)))
    return out


# ---------------------------------------------------------------- Yahoo Finance
def yahoo(symbol, inverse=False):
    url = f"https://query2.finance.yahoo.com/v8/finance/chart/{requests.utils.quote(symbol)}"
    r = _get(url, params={"period1": "-2208988800", "period2": str(int(time.time()) + 86400), "interval": "1d",
                          "includePrePost": "false", "events": "div,splits"},
             headers={"Accept": "application/json"})
    try:
        res = r.json()["chart"]["result"][0]
    except Exception as e:  # noqa: BLE001
        raise SourceError(f"yahoo bad json for {symbol}: {e}")
    ts = res.get("timestamp") or []
    q = res["indicators"]["quote"][0]
    tzname = res.get("meta", {}).get("exchangeTimezoneName") or "UTC"
    try:
        from zoneinfo import ZoneInfo
        tz = ZoneInfo(tzname)
    except Exception:  # noqa: BLE001
        tz = timezone.utc
    rows = []
    for i, t in enumerate(ts):
        o, h, l, c = q["open"][i], q["high"][i], q["low"][i], q["close"][i]
        v = (q.get("volume") or [None] * len(ts))[i]
        if c is None or o is None or h is None or l is None:
            continue
        d = datetime.fromtimestamp(t, tz).date().isoformat()
        if inverse:
            o, h, l, c = 1 / o, 1 / l, 1 / h, 1 / c
        rows.append((d, float(o), float(h), float(l), float(c), float(v or 0)))
    if len(rows) < 30:
        raise SourceError(f"yahoo returned {len(rows)} rows for {symbol}")
    rows.sort()
    return rows


def yahoo_inv(symbol):
    return yahoo(symbol, inverse=True)


# ---------------------------------------------------------------- Stooq
def stooq(symbol):
    r = _get("https://stooq.com/q/d/l/", params={"s": symbol, "i": "d"})
    txt = r.text
    if "Exceeded" in txt[:200] or "Date" not in txt[:200]:
        raise SourceError(f"stooq: {txt[:80]!r}")
    rows = []
    for rec in csv.DictReader(io.StringIO(txt)):
        d = norm_period(rec.get("Date", ""))
        o, h, l, c = (_num(rec.get(k)) for k in ("Open", "High", "Low", "Close"))
        v = _num(rec.get("Volume")) or 0.0
        if d and None not in (o, h, l, c):
            rows.append((d, o, h, l, c, v))
    if len(rows) < 30:
        raise SourceError(f"stooq returned {len(rows)} rows for {symbol}")
    rows.sort()
    return rows


# ---------------------------------------------------------------- SDMX-CSV helpers (ECB, Bundesbank)
def _parse_sdmx_csv(txt):
    """Parse SDMX-CSV (TIME_PERIOD/OBS_VALUE columns) or a plain 'date;value' style dump."""
    rows = []
    lines = txt.splitlines()
    # 1) SDMX-CSV with header
    for delim in (",", ";"):
        for i, line in enumerate(lines[:20]):
            if "TIME_PERIOD" in line and "OBS_VALUE" in line and delim in line:
                rdr = csv.DictReader(io.StringIO("\n".join(lines[i:])), delimiter=delim)
                for rec in rdr:
                    d = norm_period(rec.get("TIME_PERIOD", "") or "")
                    v = _num(rec.get("OBS_VALUE"))
                    if d:
                        rows.append((d, v))
                return rows
    # 2) generic: rows starting with a date
    for line in lines:
        parts = re.split(r"[;,\t]", line.strip())
        if len(parts) < 2:
            continue
        d = norm_period(parts[0].strip().strip('"'))
        if d is None:
            continue
        rows.append((d, _num(parts[1].strip().strip('"'))))
    return rows


def ecb(key):
    """key = FLOW/SERIES_KEY, optional ?startPeriod=... suffix."""
    flow, _, series = key.partition("/")
    series, _, extra = series.partition("?")
    params = {"format": "csvdata"}
    if extra:
        for kv in extra.split("&"):
            k, _, v = kv.partition("=")
            params[k] = v
    r = _get(f"https://data-api.ecb.europa.eu/service/data/{flow}/{series}", params=params,
             headers={"Accept": "text/csv"})
    return _clean_value_rows(_parse_sdmx_csv(r.text))


def bbk(key):
    """Bundesbank statistics REST API. key = FLOW/SERIES_KEY (e.g. BBSSY/D.REN.EUR.A620.000000WT0202.A)."""
    flow, _, series = key.partition("/")
    errors = []
    for url, headers, params in (
        (f"https://api.statistiken.bundesbank.de/rest/data/{flow}/{series}",
         {"Accept": "application/vnd.sdmx.data+csv;version=1.0.0"}, {"detail": "dataonly"}),
        (f"https://api.statistiken.bundesbank.de/rest/download/{flow}/{series}",
         {"Accept": "text/csv"}, {"format": "csv", "lang": "en"}),
    ):
        try:
            r = _get(url, headers=headers, params=params, retries=0)
            return _clean_value_rows(_parse_sdmx_csv(r.text))
        except SourceError as e:
            errors.append(str(e)[:160])
    raise SourceError("bbk: " + " | ".join(errors))


# ---------------------------------------------------------------- DBnomics
def dbnomics(key):
    """key = PROVIDER/DATASET/SERIES_CODE"""
    r = _get(f"https://api.db.nomics.world/v22/series/{key}", params={"observations": "1", "format": "json"},
             headers={"Accept": "application/json"})
    try:
        docs = r.json()["series"]["docs"]
        doc = docs[0]
    except Exception as e:  # noqa: BLE001
        raise SourceError(f"dbnomics bad json for {key}: {e}")
    rows = []
    for p, v in zip(doc.get("period", []), doc.get("value", [])):
        d = norm_period(str(p))
        if d:
            rows.append((d, _num(v)))
    return _clean_value_rows(rows)


def dbnomics_search(query, limit=12):
    """Series-level search (provider/dataset/series, name, last period) for probing."""
    r = _get("https://api.db.nomics.world/v22/series", params={"q": query, "limit": str(limit), "observations": "0"},
             headers={"Accept": "application/json"}, retries=0)
    out = []
    try:
        for doc in r.json()["series"]["docs"]:
            out.append((f"{doc['provider_code']}/{doc['dataset_code']}/{doc['series_code']}",
                        (doc.get("series_name") or "")[:90], doc.get("indexed_at", "")[:10]))
    except Exception as e:  # noqa: BLE001
        out.append(("search-error", str(e)[:120], ""))
    return out


# ---------------------------------------------------------------- CoinTrader.Pro (UDF feed behind hexchart.com, CoinMarketCap data)
def cointrader_history(symbol, start=1546300800):
    """Daily OHLCV rows (unix seconds) from the TradingView-UDF style feed used by hexchart.com."""
    r = _get("https://charts.cointrader.pro/api/history",
             params={"symbol": symbol, "resolution": "1D", "from": str(start), "to": str(int(time.time()) + 86400)},
             headers={"Accept": "application/json", "Referer": "https://charts.cointrader.pro/charts.html"}, retries=1)
    try:
        j = r.json()
    except ValueError as e:
        raise SourceError(f"cointrader bad json for {symbol}: {e}")
    if j.get("s") != "ok" or not j.get("t"):
        raise SourceError(f"cointrader {symbol}: {str(j)[:100]}")
    rows = []
    for t, o, h, l, c, v in zip(j["t"], j["o"], j["h"], j["l"], j["c"], j.get("v") or [0] * len(j["t"])):
        if None in (o, h, l, c):
            continue
        rows.append((int(t), float(o), float(h), float(l), float(c), float(v or 0)))
    rows.sort()
    return rows


# ---------------------------------------------------------------- World Bank "Pink Sheet" (monthly commodity prices since 1960)
PINK_SHEET_URL = "https://thedocs.worldbank.org/en/doc/5d903e848db1d1b83e0ec8f744e55570-0350012021/related/CMO-Historical-Data-Monthly.xlsx"


def worldbank_pinksheet(commodity):
    """Monthly price of one Pink Sheet commodity (e.g. 'Gold', 'Crude oil, WTI', 'Silver') as (date, value) rows."""
    import openpyxl  # noqa: PLC0415
    r = _get(PINK_SHEET_URL, headers={"Accept": "*/*"}, retries=1, timeout=90)
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb["Monthly Prices"] if "Monthly Prices" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    col = None
    for row in rows_iter:
        cells = [str(c).strip() if c is not None else "" for c in row]
        for i, c in enumerate(cells):
            if c.lower() == commodity.lower():
                col = i
                break
        if col is not None:
            break
    if col is None:
        raise SourceError(f"pink sheet: column {commodity!r} not found")
    out = []
    for row in rows_iter:
        if not row or row[0] is None:
            continue
        m = re.match(r"^(\d{4})M(\d{2})$", str(row[0]).strip())
        if not m:
            continue
        v = _num(row[col]) if col < len(row) else None
        if v is not None:
            out.append((f"{m.group(1)}-{m.group(2)}-01", v))
    return _clean_value_rows(out)


# ---------------------------------------------------------------- manual CSV
def manual(key, manual_dir):
    p = manual_dir / f"{key}.csv"
    if not p.exists():
        raise SourceError(f"no manual file {p.name}")
    rows = []
    for rec in csv.reader(p.open(encoding="utf-8")):
        if len(rec) < 2:
            continue
        d = norm_period(rec[0])
        if d:
            rows.append((d, _num(rec[1])))
    return _clean_value_rows(rows)


# ---------------------------------------------------------------- PBoC (best effort)
def pboc(_key):
    """Best-effort scrape of the PBoC English site 'Balance Sheet of Monetary Authority'.
    Site structure changes every year; if it breaks, put data/manual/CNCBBS.csv in place."""
    base = "http://www.pbc.gov.cn"
    idx = _get(f"{base}/en/3688247/3688975/index.html", retries=0)
    links = re.findall(r'href="([^"]+)"[^>]*>([^<]*Balance Sheet of Monetary Authority[^<]*)<', idx.text, flags=re.I)
    if not links:
        # try the "Statistics" landing page which lists yearly data sets
        idx = _get(f"{base}/en/3688247/index.html", retries=0)
        links = re.findall(r'href="([^"]+)"[^>]*>([^<]*Monetary Authority[^<]*)<', idx.text, flags=re.I)
    if not links:
        raise SourceError("pboc: no 'Balance Sheet of Monetary Authority' link found")
    href = links[0][0]
    if href.startswith("/"):
        href = base + href
    page = _get(href, retries=0)
    # find the html table with 'Total Assets' row; header cells hold months like 2025.01
    months = re.findall(r"(20\d{2})[.\-/年](\d{1,2})", page.text)
    row = re.search(r"Total\s*Assets.*?</tr>", page.text, flags=re.I | re.S)
    if not row or not months:
        raise SourceError("pboc: table layout not recognised")
    nums = [_num(x) for x in re.findall(r">\s*([\d,\.]+)\s*<", row.group(0))]
    nums = [n for n in nums if n is not None]
    if len(nums) < 1:
        raise SourceError("pboc: no numbers in Total Assets row")
    seen, periods = set(), []
    for y, m in months:
        key = f"{y}-{int(m):02d}-01"
        if key not in seen:
            seen.add(key)
            periods.append(key)
    n = min(len(nums), len(periods))
    return _clean_value_rows(list(zip(periods[:n], nums[:n])))


# ---------------------------------------------------------------- GeckoTerminal
GT = "https://api.geckoterminal.com/api/v2"
_gt_last = [0.0]


class GTHistoryLimit(SourceError):
    """Public API refuses data older than its free window (HTTP 401)."""


def _gt_get(path, params=None):
    # free tier: 30 req/min -> keep >= 2.6 s between calls; back off on 429
    for attempt in range(4):
        wait = 2.6 - (time.time() - _gt_last[0])
        if wait > 0:
            time.sleep(wait)
        h = {"User-Agent": UA, "Accept": "application/json;version=20230302"}
        try:
            r = requests.get(f"{GT}{path}", headers=h, params=params, timeout=TIMEOUT)
        except requests.RequestException as e:
            _gt_last[0] = time.time()
            if attempt == 3:
                raise SourceError(f"{path} -> {e!r}")
            time.sleep(5)
            continue
        _gt_last[0] = time.time()
        if r.status_code == 200:
            return r.json()
        if r.status_code == 429:
            time.sleep(20 * (attempt + 1))
            continue
        if r.status_code == 401:
            raise GTHistoryLimit(f"{path} -> 401 (beyond the public API history window)")
        raise SourceError(f"{path} -> HTTP {r.status_code}: {r.text[:120]!r}")
    raise SourceError(f"{path} -> rate limited (429) repeatedly")


def gt_pool_info(network, pool):
    """name, base/quote symbols and token addresses of a pool."""
    js = _gt_get(f"/networks/{network}/pools/{pool}")
    p = js.get("data", {})
    a = p.get("attributes", {})
    rel = p.get("relationships", {})
    base_id = (rel.get("base_token", {}).get("data") or {}).get("id", "")
    quote_id = (rel.get("quote_token", {}).get("data") or {}).get("id", "")
    name = a.get("name", "")
    base_sym, _, quote_sym = name.partition(" / ")
    return dict(pool=pool, name=name, base_symbol=base_sym.strip(), quote_symbol=quote_sym.split(" ")[0].strip(),
                base_token=base_id.split("_", 1)[1].lower() if "_" in base_id else "",
                quote_token=quote_id.split("_", 1)[1].lower() if "_" in quote_id else "",
                reserve_usd=float(a.get("reserve_in_usd") or 0), price_usd=_num(a.get("base_token_price_usd")))


def gt_search_pool(network, symbol, quotes, token=None):
    """Find the most liquid pool whose base token symbol matches (and token address, if given)."""
    js = _gt_get("/search/pools", params={"query": symbol, "network": network, "page": "1"})
    best, cands = None, []
    for p in js.get("data", []):
        a = p.get("attributes", {})
        name = a.get("name", "")
        base_sym, _, quote_sym = name.partition(" / ")
        base_sym = base_sym.strip()
        quote_sym = quote_sym.split(" ")[0].strip()
        rel = p.get("relationships", {})
        base_id = (rel.get("base_token", {}).get("data") or {}).get("id", "")
        base_addr = base_id.split("_", 1)[1].lower() if "_" in base_id else ""
        reserve = float(a.get("reserve_in_usd") or 0)
        cands.append((name, a.get("address"), reserve, base_addr))
        if base_sym.upper() != symbol.upper():
            continue
        if token and base_addr and base_addr != token.lower():
            continue
        score = reserve * (3.0 if quote_sym.upper() in [q.upper() for q in quotes] else 1.0)
        if best is None or score > best[0]:
            best = (score, dict(pool=a.get("address"), name=name, reserve_usd=reserve,
                                base_symbol=base_sym, quote_symbol=quote_sym, base_token=base_addr,
                                price_usd=_num(a.get("base_token_price_usd"))))
    return (best[1] if best else None), cands


def gt_ohlcv(network, pool, timeframe="day", aggregate=1, limit=1000, before=None, token="base"):
    params = {"aggregate": str(aggregate), "limit": str(limit), "currency": "usd", "token": token}
    if before:
        params["before_timestamp"] = str(int(before))
    js = _gt_get(f"/networks/{network}/pools/{pool}/ohlcv/{timeframe}", params=params)
    lst = js.get("data", {}).get("attributes", {}).get("ohlcv_list", []) or []
    rows = []
    for t, o, h, l, c, v in lst:
        if None in (o, h, l, c):
            continue
        rows.append((int(t), float(o), float(h), float(l), float(c), float(v or 0)))
    rows.sort()
    return rows


def gt_ohlcv_history(network, pool, timeframe, aggregate, pages=1, token="base"):
    """Page backwards until the API returns nothing, refuses (401 = free window exhausted) or `pages` is hit."""
    allrows, before = [], None
    for _ in range(pages):
        try:
            rows = gt_ohlcv(network, pool, timeframe, aggregate, before=before, token=token)
        except GTHistoryLimit:
            break
        if len(rows) < 2:
            break
        allrows = rows + allrows
        before = rows[0][0] - 1
    seen, out = set(), []
    for r in allrows:
        if r[0] not in seen:
            seen.add(r[0])
            out.append(r)
    out.sort()
    return out
